"""Read a folder of documents and split them into searchable chunks.

Supported: .pdf, .docx, .txt, .md, .csv
Every chunk remembers which file (and PDF page) it came from, so answers can cite sources.
"""
from __future__ import annotations

import os
import re
from dataclasses import dataclass, asdict

SUPPORTED = {".pdf", ".docx", ".txt", ".md", ".csv"}


@dataclass
class Chunk:
    text: str
    source: str        # file path relative to the folder
    page: int | None   # PDF page number (1-based), None for other types
    chunk_id: int

    def to_dict(self):
        return asdict(self)

    def label(self) -> str:
        return f"{self.source} p.{self.page}" if self.page else self.source


def _read_pdf(path):
    from pypdf import PdfReader
    reader = PdfReader(path)
    for i, page in enumerate(reader.pages, start=1):
        yield i, page.extract_text() or ""


def _read_docx(path):
    import docx
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    yield None, "\n".join(parts)


def _read_text(path):
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        yield None, f.read()


READERS = {".pdf": _read_pdf, ".docx": _read_docx}


def split_text(text: str, size: int = 900, overlap: int = 150) -> list[str]:
    """Split text into ~size-character chunks on paragraph/sentence boundaries."""
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)       # join words hyphenated across lines
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        return []
    # break into sentence-ish pieces, then pack them into chunks
    pieces = re.split(r"(?<=[.!?])\s+|\n\n", text)
    chunks, cur = [], ""
    for p in pieces:
        p = p.strip()
        if not p:
            continue
        while len(p) > size:  # very long piece with no punctuation
            if cur:
                chunks.append(cur.strip())
                cur = ""
            chunks.append(p[:size])
            p = p[size - overlap:]
        if len(cur) + len(p) + 1 > size and cur:
            chunks.append(cur.strip())
            cur = cur[-overlap:] + " " + p  # carry overlap for context
        else:
            cur = (cur + " " + p) if cur else p
    if cur.strip():
        chunks.append(cur.strip())
    return chunks


def list_files(folder: str) -> list[str]:
    out = []
    for root, dirs, files in os.walk(folder):
        dirs[:] = [d for d in dirs if not d.startswith(".") and d != "__pycache__"]
        for name in sorted(files):
            if os.path.splitext(name)[1].lower() in SUPPORTED and not name.startswith("~$"):
                out.append(os.path.join(root, name))
    return sorted(out)


def load_folder(folder: str, progress=None) -> tuple[list[Chunk], list[str]]:
    """Return (chunks, errors). `progress(i, total, filename)` is optional."""
    files = list_files(folder)
    chunks, errors = [], []
    for i, path in enumerate(files):
        rel = os.path.relpath(path, folder)
        if progress:
            progress(i, len(files), rel)
        ext = os.path.splitext(path)[1].lower()
        reader = READERS.get(ext, _read_text)
        try:
            for page, text in reader(path):
                for piece in split_text(text):
                    chunks.append(Chunk(piece, rel, page, len(chunks)))
        except Exception as e:  # keep going if one file is broken
            errors.append(f"{rel}: {e}")
    return chunks, errors
